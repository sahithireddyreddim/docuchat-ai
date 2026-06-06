"""
Document processor - extracts text from PDF, DOCX, and TXT files.
OCR is optional - only used if pytesseract is installed.
"""

import os
import logging
from typing import Optional
import PyPDF2
import docx
from langchain_text_splitters import RecursiveCharacterTextSplitter

from config import settings

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF. Tries OCR if no text found and pytesseract is available."""
    text = ""
    try:
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page_num, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    text += f"\n[Page {page_num + 1}]\n{page_text}"
    except Exception as e:
        logger.warning(f"PDF extraction failed: {e}")

    # Try OCR only if pytesseract is installed and no text found
    if len(text.strip()) < 50:
        try:
            import pytesseract
            from pdf2image import convert_from_path
            logger.info("Attempting OCR...")
            images = convert_from_path(file_path, dpi=200)
            for page_num, image in enumerate(images):
                page_text = pytesseract.image_to_string(image, lang="eng")
                text += f"\n[Page {page_num + 1} - OCR]\n{page_text}"
        except ImportError:
            logger.info("pytesseract not installed - skipping OCR")
        except Exception as e:
            logger.error(f"OCR failed: {e}")

    return text.strip()


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from DOCX."""
    doc = docx.Document(file_path)
    paragraphs = []

    for para in doc.paragraphs:
        if para.text.strip():
            if para.style.name.startswith("Heading"):
                paragraphs.append(f"\n## {para.text}")
            else:
                paragraphs.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    return "\n".join(paragraphs)


def extract_text_from_txt(file_path: str) -> str:
    """Read plain text file."""
    for encoding in ["utf-8", "latin-1", "cp1252"]:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    return "[Could not decode text file]"


def extract_text(file_path: str, file_type: str) -> str:
    extractors = {
        ".pdf": extract_text_from_pdf,
        ".docx": extract_text_from_docx,
        ".txt": extract_text_from_txt,
    }
    extractor = extractors.get(file_type)
    if not extractor:
        raise ValueError(f"Unsupported file type: {file_type}")
    return extractor(file_path)


def split_into_chunks(text: str, document_name: str) -> list[dict]:
    """Split text into overlapping chunks for RAG."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE,
        chunk_overlap=settings.CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    chunks = splitter.split_text(text)

    return [
        {
            "text": chunk,
            "metadata": {
                "document": document_name,
                "chunk_index": i,
                "total_chunks": len(chunks),
            },
        }
        for i, chunk in enumerate(chunks)
    ]
